from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml
from docx import Document

from agent_sdk.context_selection import DesignStage, TaskAwareContextSelector
from agent_sdk.specification_gate import (
    Gate1ArtifactStore,
    RequirementEntry,
    SpecificationGate,
    UnifiedSpecification,
    VersionMetadata,
    classify_version_change,
)
from agent_sdk.specifications import (
    DocumentFormat,
    ScriptedVisionAdapter,
    SourceRef,
    SpecificationCategory,
    SpecificationDocument,
    SpecificationPreprocessor,
    VisionProposal,
    VisionStatus,
)


def _write_manifest(root: Path) -> None:
    (root / "functional").mkdir()
    (root / "architecture").mkdir()
    (root / "interfaces").mkdir()
    (root / "pdk").mkdir()
    (root / "functional/requirements.md").write_text(
        "# Accumulator\nThe accumulator exposes a ready signal.\n", encoding="utf-8"
    )
    (root / "architecture/clocking.txt").write_text(
        "The accumulator uses one synchronous clock domain.\n", encoding="utf-8"
    )
    (root / "interfaces/ready.png").write_bytes(b"not-a-real-png-but-a-source-asset")
    (root / "pdk/library.txt").write_text("Clock library corner constraints.\n", encoding="utf-8")
    manifest = {
        "functional_spec": {
            "documents": [
                {
                    "id": "REQ-FUNC-001",
                    "title": "Functional requirements",
                    "format": "md",
                    "path": "functional/requirements.md",
                }
            ]
        },
        "architectural_spec": {
            "documents": [
                {
                    "id": "REQ-ARCH-001",
                    "title": "Clocking architecture",
                    "format": "txt",
                    "path": "architecture/clocking.txt",
                }
            ]
        },
        "interface_spec": {
            "documents": [
                {
                    "id": "REQ-IFACE-001",
                    "title": "Ready waveform",
                    "format": "png",
                    "path": "interfaces/ready.png",
                }
            ]
        },
        "pdk_spec": {
            "documents": [
                {
                    "id": "REQ-PDK-001",
                    "title": "PDK library assumptions",
                    "format": "txt",
                    "path": "pdk/library.txt",
                }
            ]
        },
    }
    (root / "specification-manifest.yaml").write_text(yaml.safe_dump(manifest), encoding="utf-8")


def test_preprocessor_preserves_source_references_and_legacy_manifest_tree(tmp_path: Path):
    _write_manifest(tmp_path)
    processor = SpecificationPreprocessor(tmp_path)

    manifest = processor.load_manifest()
    trees = processor.process_manifest(manifest)
    target = processor.persist_tree(trees[0])

    assert [tree.document_id for tree in trees] == [
        "REQ-FUNC-001",
        "REQ-ARCH-001",
        "REQ-IFACE-001",
        "REQ-PDK-001",
    ]
    assert trees[0].nodes[0].source.document_id == "REQ-FUNC-001"
    assert trees[0].nodes[0].source.location == "line:1"
    assert target.is_file()
    assert yaml.safe_load(target.read_text())["source_hash"] == trees[0].source_hash


@pytest.mark.anyio
async def test_image_resolution_retries_typed_proposals_then_marks_review_or_accepts(
    tmp_path: Path,
):
    _write_manifest(tmp_path)
    processor = SpecificationPreprocessor(tmp_path)
    image_tree = next(
        tree
        for tree in processor.process_manifest(processor.load_manifest())
        if tree.format is DocumentFormat.PNG
    )

    accepted = await processor.resolve_images(
        image_tree,
        ScriptedVisionAdapter(
            [
                VisionProposal(confidence=0.2, structure={"bad": True}),
                VisionProposal(
                    confidence=0.9, structure={"kind": "waveform", "signals": ["ready"]}
                ),
            ]
        ),
    )
    review = await processor.resolve_images(image_tree)

    assert accepted.nodes[0].vision_status is VisionStatus.ACCEPTED
    assert accepted.nodes[0].resolved_structure == {"kind": "waveform", "signals": ["ready"]}
    assert review.nodes[0].vision_status is VisionStatus.REVIEW_REQUIRED


def test_task_aware_selection_prunes_categories_before_keyword_matching(tmp_path: Path):
    _write_manifest(tmp_path)
    processor = SpecificationPreprocessor(tmp_path)
    trees = processor.process_manifest(processor.load_manifest())

    selected = TaskAwareContextSelector().select(
        trees,
        DesignStage.RTL_DEVELOPMENT,
        "Implement accumulator ready clock logic",
        ["line:2"],
    )

    assert "REQ-FUNC-001" in selected.selected_document_ids
    assert "REQ-ARCH-001" in selected.selected_document_ids
    assert "REQ-PDK-001" not in selected.selected_document_ids


def test_docx_and_vsdx_parsers_preserve_structured_source_locations(tmp_path: Path):
    docx = Document()
    docx.add_paragraph("Clocking architecture")
    table = docx.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "signal"
    table.rows[0].cells[1].text = "ready"
    docx.save(tmp_path / "architecture.docx")
    with ZipFile(tmp_path / "topology.vsdx", "w") as archive:
        archive.writestr("visio/pages/page1.xml", "<Page><Shape Name='accumulator'/></Page>")

    processor = SpecificationPreprocessor(tmp_path)
    docx_tree = processor.process_document(
        SpecificationDocument(
            id="REQ-ARCH-001",
            title="Architecture",
            format="docx",
            path="architecture.docx",
            category="architectural",
        )
    )
    vsdx_tree = processor.process_document(
        SpecificationDocument(
            id="REQ-ARCH-002",
            title="Topology",
            format="vsdx",
            path="topology.vsdx",
            category="architectural",
        )
    )

    assert [node.source.location for node in docx_tree.nodes] == ["paragraph:1", "table:1"]
    assert vsdx_tree.nodes[0].source.location == "vsdx-part:visio/pages/page1.xml"


def test_gate_one_detects_traceability_and_verifiability_then_persists_soft_lock(tmp_path: Path):
    _write_manifest(tmp_path)
    processor = SpecificationPreprocessor(tmp_path)
    trees = processor.process_manifest(processor.load_manifest())
    source = SourceRef(
        document_id="REQ-FUNC-001",
        relative_path="functional/requirements.md",
        source_hash=trees[0].source_hash,
        format=DocumentFormat.MD,
        location="line:2",
    )
    specification = UnifiedSpecification(
        version="1.0.0",
        documents=trees,
        requirements=[
            RequirementEntry(
                id="REQ-ACC-001",
                category=SpecificationCategory.FUNCTIONAL,
                text="Accumulator must expose ready.",
                source_refs=[source],
                dependencies=["REQ-MISSING-001"],
            )
        ],
    )
    gate = SpecificationGate()
    graph, report = gate.validate(
        specification,
        required_categories={SpecificationCategory.FUNCTIONAL, SpecificationCategory.ARCHITECTURAL},
    )
    kind, rationale = classify_version_change(None, specification)
    metadata = VersionMetadata(
        version="1.0.0",
        change_kind=kind,
        rationale=rationale,
        unified_specification_hash=trees[0].source_hash,
    )

    rejected = gate.soft_lock(specification, report, metadata, user_approved=True)
    accepted = gate.soft_lock(
        specification,
        report,
        metadata,
        user_approved=True,
        proceed_with_gaps=True,
    )
    Gate1ArtifactStore(tmp_path).persist(
        specification, graph, report, accepted.metadata, plans=[{"plan_id": "gap-plan"}]
    )

    assert rejected.accepted is False
    assert accepted.accepted is True
    assert accepted.metadata and accepted.metadata.soft_locked
    assert {gap.type.value for gap in report.gaps} == {"traceability", "verifiability"}
    for relative in [
        "unified-specification.yaml",
        "dependency-graph.yaml",
        "gap-report.yaml",
        "version-metadata.yaml",
        "plans/plan-1.yaml",
    ]:
        assert (tmp_path / relative).is_file()


def test_gate_one_counts_transitive_dependents_for_missing_requirement():
    source = SourceRef(
        document_id="REQ-FUNC-001",
        relative_path="functional/requirements.md",
        source_hash="source-hash",
        format=DocumentFormat.MD,
        location="line:1",
    )
    specification = UnifiedSpecification(
        version="1.0.0",
        documents=[],
        requirements=[
            RequirementEntry(
                id="REQ-A",
                category=SpecificationCategory.FUNCTIONAL,
                text="A depends on an undefined requirement.",
                source_refs=[source],
                dependencies=["REQ-MISSING"],
                acceptance_checks=["test-a"],
            ),
            RequirementEntry(
                id="REQ-B",
                category=SpecificationCategory.FUNCTIONAL,
                text="B depends on A.",
                source_refs=[source],
                dependencies=["REQ-A"],
                acceptance_checks=["test-b"],
            ),
            RequirementEntry(
                id="REQ-C",
                category=SpecificationCategory.FUNCTIONAL,
                text="C depends on B.",
                source_refs=[source],
                dependencies=["REQ-B"],
                acceptance_checks=["test-c"],
            ),
        ],
    )

    _graph, report = SpecificationGate().validate(specification, required_categories=set())

    missing_dependency = next(
        gap for gap in report.gaps if gap.locations == ["REQ-A", "REQ-MISSING"]
    )
    assert missing_dependency.blast_radius == 3


def test_gate_one_version_preview_marks_existing_requirement_change_major():
    source = SourceRef(
        document_id="REQ-FUNC-001",
        relative_path="functional/requirements.md",
        source_hash="source-hash",
        format=DocumentFormat.MD,
        location="line:1",
    )
    previous = UnifiedSpecification(
        version="1.0.0",
        documents=[],
        requirements=[
            RequirementEntry(
                id="REQ-FUNC-001",
                category=SpecificationCategory.FUNCTIONAL,
                text="Support one accumulator lane.",
                source_refs=[source],
            )
        ],
    )
    current = previous.model_copy(
        update={
            "version": "2.0.0",
            "requirements": [
                previous.requirements[0].model_copy(
                    update={"text": "Support two accumulator lanes."}
                )
            ],
        }
    )

    change_kind, rationale = classify_version_change(previous, current)

    assert change_kind.value == "major"
    assert rationale == ["Existing requirement changed: REQ-FUNC-001 (text)."]
